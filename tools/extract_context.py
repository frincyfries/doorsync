from pathlib import Path
import json

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BRD = Path(r"C:\Users\Aithe\Downloads\DoorSync_BRD.docx")
DB = Path(r"C:\Users\Aithe\Downloads\DoorSync_Database_Design_withERD.xlsx")
OUT = ROOT / "docs" / "source-summary"


def extract_docx(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "style": paragraph.style.name if paragraph.style else "",
                    "text": text,
                }
            )

    tables = []
    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    return {"paragraphs": paragraphs, "tables": tables}


def extract_xlsx(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False)
    sheets = {}
    for sheet in workbook.worksheets:
        used = []
        for row in sheet.iter_rows():
            values = []
            has_value = False
            for cell in row:
                value = cell.value
                if value is not None:
                    has_value = True
                values.append(value)
            if has_value:
                while values and values[-1] is None:
                    values.pop()
                used.append(values)
        sheets[sheet.title] = used
    return {"sheets": sheets}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    payload = {
        "brd": extract_docx(BRD),
        "database": extract_xlsx(DB),
    }
    (OUT / "extracted_context.json").write_text(
        json.dumps(payload, indent=2, default=str),
        encoding="utf-8",
    )

    print(f"Wrote {OUT / 'extracted_context.json'}")
    print("BRD paragraphs:", len(payload["brd"]["paragraphs"]))
    print("BRD tables:", len(payload["brd"]["tables"]))
    print("Workbook sheets:", ", ".join(payload["database"]["sheets"].keys()))


if __name__ == "__main__":
    main()
